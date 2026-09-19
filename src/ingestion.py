from pathlib import Path
from typing import List
from hashlib import sha256

from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_community.vectorstores import FAISS
from langchain_text_splitters import RecursiveCharacterTextSplitter
from docx import Document as DocxDocument

from src.config import get_settings
from src.vectorstore import get_vector_store, save_vector_store, get_embeddings

SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def _load_docx(path: Path) -> List[Document]:
    docx = DocxDocument(str(path))
    text = "\n".join(p.text for p in docx.paragraphs if p.text.strip())
    return [
        Document(
            page_content=text,
            metadata={"source": str(path), "title": path.name},
        )
    ]


def load_file(path: Path) -> List[Document]:
    ext = path.suffix.lower()

    if ext == ".pdf":
        docs = PyPDFLoader(str(path)).load()
    elif ext in [".txt", ".text"]:
        return TextLoader(str(path)).load()
    elif ext in [".md", ".markdown"]:
        return TextLoader(str(path)).load()
    elif ext == ".docx":
        docs = _load_docx(path)
    else:
        raise ValueError(
            f"Unsupported file type: {ext}. Use PDF, TXT, MD, or DOCX."
        )

    for d in docs:
        d.metadata.setdefault("source", str(path))
        d.metadata.setdefault("title", path.name)
        d.metadata["document_name"] = path.name

    return docs


def _stable_chunk_id(path: Path, chunk: Document, position: int) -> str:
    """Stable IDs make repeat ingestion idempotent instead of creating duplicates."""
    material = f"{path.name}|{position}|{chunk.page_content}".encode("utf-8")
    digest = sha256(material).hexdigest()[:24]
    safe_stem = "".join(
        c if c.isalnum() or c in "-_" else "-" for c in path.stem
    )[:60]
    return f"{safe_stem}-{position}-{digest}"


def ingest_file(path: Path) -> int:
    docs = load_file(path)
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=160,
        separators=["\n\n", "\n", ".", " ", ""],
    )
    chunks = splitter.split_documents(docs)

    for i, chunk in enumerate(chunks):
        chunk.metadata["chunk_index"] = i
        chunk.metadata["knowledge_base"] = "tunajenga"

    ids = [_stable_chunk_id(path, chunk, i) for i, chunk in enumerate(chunks)]

    store = get_vector_store()
    if store is None:
        store = FAISS.from_documents(chunks, get_embeddings(), ids=ids)
    else:
        # Delete any existing chunks with these IDs first (idempotent re-ingest)
        try:
            store.delete(ids=ids)
        except Exception:
            pass  # IDs weren't in the index yet — that's fine
        store.add_documents(chunks, ids=ids)

    save_vector_store(store)
    return len(chunks)


def ingest_directory(directory: Path) -> int:
    total = 0
    if not directory.exists():
        return 0
    for path in sorted(directory.iterdir()):
        if path.is_file() and path.suffix.lower() in SUPPORTED:
            total += ingest_file(path)
    return total


def namespace() -> str:
    return "tunajenga"






 